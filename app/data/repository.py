"""Load and query the ParcelPilot Excel dataset."""

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.data.models import Account, DatasetMetadata, Order, Ticket

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = PROJECT_ROOT / "data" / "ParcelPilot_Assessment_Data.xlsx"

DATETIME_FORMAT = "%Y-%m-%d %H:%M"


def _optional_str(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _required_str(value: object, field: str) -> str:
    text = _optional_str(value)
    if text is None:
        raise ValueError(f"Missing required value for {field}")
    return text


def _as_bool(value: object, field: str) -> bool:
    if isinstance(value, bool):
        return value
    text = _required_str(value, field).upper()
    if text == "TRUE":
        return True
    if text == "FALSE":
        return False
    raise ValueError(f"Invalid boolean for {field}: {value!r}")


def _as_int(value: object, field: str) -> int:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Invalid integer for {field}: {value!r}")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = _required_str(value, field)
    return int(text)


def _as_datetime(value: object, field: str) -> datetime:
    if isinstance(value, datetime):
        return value
    text = _required_str(value, field)
    return datetime.strptime(text, DATETIME_FORMAT)


def _as_optional_datetime(value: object, field: str) -> datetime | None:
    if _optional_str(value) is None and not isinstance(value, datetime):
        return None
    return _as_datetime(value, field)


def _header_map(worksheet: Worksheet) -> dict[str, int]:
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    mapping: dict[str, int] = {}
    for index, name in enumerate(header_row):
        if name is None:
            continue
        mapping[str(name).strip()] = index
    return mapping


def _row_dict(headers: dict[str, int], row: tuple[object, ...]) -> dict[str, object]:
    return {name: row[index] if index < len(row) else None for name, index in headers.items()}


class ParcelPilotRepository:
    """In-memory repository backed by the assessment Excel workbook."""

    def __init__(self, workbook_path: Path | None = None) -> None:
        self.workbook_path = Path(workbook_path) if workbook_path else DEFAULT_WORKBOOK_PATH
        self.metadata: DatasetMetadata
        self.accounts: list[Account] = []
        self.orders: list[Order] = []
        self.tickets: list[Ticket] = []
        self._accounts_by_id: dict[str, Account] = {}
        self._orders_by_id: dict[str, Order] = {}
        self._tickets_by_id: dict[str, Ticket] = {}
        self._load()

    def _load(self) -> None:
        workbook = load_workbook(self.workbook_path, data_only=True, read_only=True)
        try:
            self.metadata = self._load_metadata(workbook["README"])
            self.accounts = self._load_accounts(workbook["accounts"])
            self.orders = self._load_orders(workbook["orders"])
            self.tickets = self._load_tickets(workbook["tickets"])
        finally:
            workbook.close()

        self._accounts_by_id = {account.account_id: account for account in self.accounts}
        self._orders_by_id = {order.order_id: order for order in self.orders}
        self._tickets_by_id = {ticket.ticket_id: ticket for ticket in self.tickets}

    def _load_metadata(self, worksheet: Worksheet) -> DatasetMetadata:
        values: dict[str, str] = {}
        title = ""
        for row in worksheet.iter_rows(min_row=1, values_only=True):
            label = _optional_str(row[0] if row else None)
            value = _optional_str(row[1] if row and len(row) > 1 else None)
            if not label:
                continue
            if not value:
                if not title:
                    title = label
                continue
            values[label] = value
        return DatasetMetadata(
            title=title,
            snapshot=values["Dataset snapshot"],
            currency=values["Currency"],
            notes=values.get("Notes"),
            important=values.get("Important"),
        )

    def _load_accounts(self, worksheet: Worksheet) -> list[Account]:
        headers = _header_map(worksheet)
        accounts: list[Account] = []
        for raw in worksheet.iter_rows(min_row=2, values_only=True):
            row = _row_dict(headers, raw)
            if _optional_str(row.get("account_id")) is None:
                continue
            accounts.append(
                Account(
                    account_id=_required_str(row.get("account_id"), "account_id"),
                    account_name=_required_str(row.get("account_name"), "account_name"),
                    plan=_required_str(row.get("plan"), "plan"),
                    status=_required_str(row.get("status"), "status"),
                    csm=_required_str(row.get("csm"), "csm"),
                    contract_file=_optional_str(row.get("contract_file")),
                    premium_support=_as_bool(row.get("premium_support"), "premium_support"),
                    notes=_optional_str(row.get("notes")),
                )
            )
        return accounts

    def _load_orders(self, worksheet: Worksheet) -> list[Order]:
        headers = _header_map(worksheet)
        orders: list[Order] = []
        for raw in worksheet.iter_rows(min_row=2, values_only=True):
            row = _row_dict(headers, raw)
            if _optional_str(row.get("order_id")) is None:
                continue
            orders.append(
                Order(
                    order_id=_required_str(row.get("order_id"), "order_id"),
                    account_id=_required_str(row.get("account_id"), "account_id"),
                    carrier=_required_str(row.get("carrier"), "carrier"),
                    status=_required_str(row.get("status"), "status"),
                    booked_at=_as_datetime(row.get("booked_at"), "booked_at"),
                    pickup_window_start=_as_datetime(
                        row.get("pickup_window_start"), "pickup_window_start"
                    ),
                    pickup_window_end=_as_datetime(
                        row.get("pickup_window_end"), "pickup_window_end"
                    ),
                    pickup_actual_at=_as_optional_datetime(
                        row.get("pickup_actual_at"), "pickup_actual_at"
                    ),
                    shipment_fee_inr=_as_int(row.get("shipment_fee_inr"), "shipment_fee_inr"),
                    carrier_fault=_as_bool(row.get("carrier_fault"), "carrier_fault"),
                    customer_fault=_as_bool(row.get("customer_fault"), "customer_fault"),
                    cancellation_requested_at=_as_optional_datetime(
                        row.get("cancellation_requested_at"), "cancellation_requested_at"
                    ),
                    notes=_optional_str(row.get("notes")),
                )
            )
        return orders

    def _load_tickets(self, worksheet: Worksheet) -> list[Ticket]:
        headers = _header_map(worksheet)
        tickets: list[Ticket] = []
        for raw in worksheet.iter_rows(min_row=2, values_only=True):
            row = _row_dict(headers, raw)
            if _optional_str(row.get("ticket_id")) is None:
                continue
            tickets.append(
                Ticket(
                    ticket_id=_required_str(row.get("ticket_id"), "ticket_id"),
                    account_id=_required_str(row.get("account_id"), "account_id"),
                    created_at=_as_datetime(row.get("created_at"), "created_at"),
                    status=_required_str(row.get("status"), "status"),
                    subject=_required_str(row.get("subject"), "subject"),
                    description=_required_str(row.get("description"), "description"),
                    channel=_required_str(row.get("channel"), "channel"),
                    assigned_to=_required_str(row.get("assigned_to"), "assigned_to"),
                    last_customer_message_at=_as_datetime(
                        row.get("last_customer_message_at"), "last_customer_message_at"
                    ),
                    historical_resolution=_optional_str(row.get("historical_resolution")),
                )
            )
        return tickets

    def get_account(self, account_id: str) -> Account | None:
        return self._accounts_by_id.get(account_id)

    def get_account_by_name(self, account_name: str) -> Account | None:
        needle = account_name.strip().casefold()
        matches = [
            account
            for account in self.accounts
            if account.account_name.casefold() == needle
        ]
        if len(matches) != 1:
            return None
        return matches[0]

    def list_accounts(self) -> list[Account]:
        return list(self.accounts)

    def get_order(self, order_id: str) -> Order | None:
        return self._orders_by_id.get(order_id)

    def list_orders(self, account_id: str | None = None) -> list[Order]:
        if account_id is None:
            return list(self.orders)
        return [order for order in self.orders if order.account_id == account_id]

    def get_ticket(self, ticket_id: str) -> Ticket | None:
        return self._tickets_by_id.get(ticket_id)

    def list_tickets(
        self,
        account_id: str | None = None,
        status: str | None = None,
    ) -> list[Ticket]:
        tickets = self.tickets
        if account_id is not None:
            tickets = [ticket for ticket in tickets if ticket.account_id == account_id]
        if status is not None:
            tickets = [ticket for ticket in tickets if ticket.status.value == status]
        return list(tickets)
